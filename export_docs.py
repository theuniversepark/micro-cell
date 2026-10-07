"""docs/*.html → docs/*.pdf (Chrome 인쇄, A3 가로) · docs/*.docx (python-docx로 직접 생성, A3 가로).
GitHub Pages의 PDF·Word 저장 버튼용. Chrome이 없으면 건너뜀.
실행: uv run --with python-docx --with beautifulsoup4 python export_docs.py"""
import os
import re
import subprocess
import sys
import tempfile

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PAGES = ["precision_zone.html", "ocs_precision.html", "dmworks_zones.html"]
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
FONT = "맑은 고딕"
NAVY, NAVY2, TEAL, INK, MUTED = "10214B", "1B3A73", "0090A0", "13203F", "56637F"
PAGE_W, PAGE_H, MARGIN = Mm(420), Mm(297), Mm(12)
CONTENT_W = PAGE_W - 2 * MARGIN
FIG_PX = 2000


# ------------------------------------------------------------------ Chrome
def chrome(*args):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", *args],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False, timeout=180)


def pdf(src, dst):
    chrome("--no-pdf-header-footer", f"--print-to-pdf={dst}", "file://" + os.path.abspath(src))


def svg_png(svg, css, tmp, n):
    vb = re.search(r'viewBox="([-\d.]+) ([-\d.]+) ([\d.]+) ([\d.]+)"', svg, re.I)
    w, h = (float(vb.group(3)), float(vb.group(4))) if vb else (1000.0, 600.0)
    H = max(40, round(FIG_PX * h / w))
    svg = re.sub(r'style="max-width:[^"]*"', "", svg, count=1)
    page, shot = os.path.join(tmp, f"f{n}.html"), os.path.join(tmp, f"f{n}.png")
    open(page, "w", encoding="utf-8").write(
        f'<!doctype html><meta charset="utf-8">{css}<style>:root{{color-scheme:light}}html,body{{margin:0;background:#fff}}'
        f'svg{{display:block;width:{FIG_PX}px;height:{H}px;max-width:none!important}}</style>{svg}')
    chrome(f"--window-size={FIG_PX},{H}", "--force-device-scale-factor=1", f"--screenshot={shot}", "file://" + page)
    return (shot if os.path.exists(shot) else None), w / h


# ------------------------------------------------------------------ docx helpers
def set_font(run, size=None, bold=None, color=None, italic=None, mono=False):
    run.font.name = "Consolas" if mono else FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    sh = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        sh.set(qn(k), v)
    cell._element.get_or_add_tcPr().append(sh)


def table_props(table, border="C9D3E3"):
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", border)):
            e.set(qn(k), v)
        b.append(e)
    tblpr.append(b)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", 40), ("left", 60), ("bottom", 40), ("right", 60)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def row_flag(row, tag):
    e = OxmlElement(f"w:{tag}")
    if tag == "tblHeader":
        e.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(e)


def hyperlink(par, url, text, size, color=TEAL):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    rpr.append(fonts)
    for tag, val in (("w:color", color), ("w:u", "single"), ("w:sz", str(int(size * 2)))):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def tight(par, before=0, after=0, line=1.0):
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


# ------------------------------------------------------------------ HTML → docx
def clean(soup):
    """화면 전용 요소(이동 바·목차·'도면 ↑'·'탑재 SW ↓' 안내·막대그래프)는 문서에서 뺌."""
    for sel in ["nav", "script", "span.where", 'a.tolink[href^="#topo-"]:not(.sw)', "span.meter"]:
        for e in soup.select(sel):
            e.decompose()


def inline(par, node, size, bold=False, color=INK, mono=False):
    for ch in node.children:
        if isinstance(ch, NavigableString):
            txt = re.sub(r"\s+", " ", str(ch))
            if txt.strip():
                set_font(par.add_run(txt), size, bold, color, mono=mono)
            elif txt == " " and par.runs:
                set_font(par.add_run(" "), size)
        elif isinstance(ch, Tag):
            if ch.name == "br":
                par.add_run().add_break()
            elif ch.name == "a" and ch.get("href", "").startswith("http"):
                hyperlink(par, ch["href"], ch.get_text(" ", strip=True), size)
                par.add_run(" ")
            elif ch.name in ("b", "strong"):
                inline(par, ch, size, True, color, mono)
            elif ch.name == "code":
                inline(par, ch, size, bold, TEAL, True)
            elif ch.name != "svg":
                inline(par, ch, size, bold, color, mono)


def text_mm(txt, fs):
    """글자 폭 추정(mm): 한글·전각 1.0em, 그 외 0.56em."""
    em = fs * 0.3528
    return sum(em * (1.0 if ord(c) > 0x2E80 else 0.56) for c in txt)


def col_widths(grid, ncol, fs):
    """짧은 열(기호·수량·단위·단가·금액 등)은 글자 폭만큼 고정해 줄바꿈을 막고, 나머지 폭은 긴 열에 글자 수 비례로 나눔."""
    total_mm = CONTENT_W / Mm(1)
    maxlen, need = [0] * ncol, [0.0] * ncol
    for r in grid:
        for i, (txt, span) in enumerate(r):
            if span == 1:
                maxlen[i] = max(maxlen[i], len(txt))
                need[i] = max(need[i], text_mm(txt, fs) + 3.0)
    fixed = {i: min(need[i], 42.0) for i in range(ncol) if maxlen[i] <= 14}
    rest = total_mm - sum(fixed.values())
    flex = [i for i in range(ncol) if i not in fixed]
    if flex:
        wts = {i: max(min(maxlen[i], 90), 8) ** 0.8 for i in flex}
        tw = sum(wts.values())
        floor = 14.0
        alloc = {i: max(floor, rest * wts[i] / tw) for i in flex}
        scale = rest / sum(alloc.values())
        alloc = {i: v * scale for i, v in alloc.items()}
    else:
        alloc, k = {}, total_mm / sum(fixed.values())
        fixed = {i: v * k for i, v in fixed.items()}
    return [Mm(fixed.get(i, alloc.get(i, 10))) for i in range(ncol)]


class Writer:
    def __init__(self, css, tmp):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = PAGE_W, PAGE_H
        for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, k, MARGIN)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(9)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        self.css, self.tmp, self.nfig, self.first = css, tmp, 0, True

    def heading(self, el, size, color, before):
        p = self.doc.add_paragraph()
        tight(p, before, 4)
        p.paragraph_format.keep_with_next = True
        inline(p, el, size, True, color)

    def para(self, el, size=9, color=INK, box=False, italic=False):
        if not el.get_text(strip=True):
            return
        p = self.doc.add_paragraph()
        tight(p, 2, 4, 1.15)
        inline(p, el, size, False, color)
        if italic:
            for r in p.runs:
                r.font.italic = True
        if box:
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:color", "E07A1F"), ("w:space", "6")):
                left.set(qn(k), v)
            bd.append(left)
            ppr.append(bd)
            sh = OxmlElement("w:shd")
            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "FFF6EC")):
                sh.set(qn(k), v)
            ppr.append(sh)

    def bullets(self, el, size=9):
        for li in el.find_all("li", recursive=False):
            p = self.doc.add_paragraph(style="List Bullet")
            tight(p, 0, 2, 1.1)
            inline(p, li, size)

    def figure_svg(self, svg, max_h_mm=215):
        if not svg.get("viewbox") and not svg.get("viewBox"):
            return
        shot, ratio = svg_png(str(svg), self.css, self.tmp, self.nfig)
        self.nfig += 1
        if not shot:
            return
        w = CONTENT_W
        if w / ratio > Mm(max_h_mm):
            w = int(Mm(max_h_mm) * ratio)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 4, 4)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(shot, width=w)

    def legend(self, ul):
        items = [li.get_text(" ", strip=True) for li in ul.find_all("li")]
        if items:
            p = self.doc.add_paragraph()
            tight(p, 0, 4)
            set_font(p.add_run("범례  "), 8, True, NAVY2)
            set_font(p.add_run("  ·  ".join(items)), 8, False, MUTED)

    def table(self, tbl):
        rows = []
        for sect, tag in (("head", "thead"), ("body", "tbody"), ("foot", "tfoot")):
            for part in tbl.find_all(tag, recursive=False):
                rows += [(sect, tr) for tr in part.find_all("tr", recursive=False)]
        if not rows:
            rows = [("body", tr) for tr in tbl.find_all("tr")]
        grid = []
        for _, tr in rows:
            cells = []
            for c in tr.find_all(["td", "th"], recursive=False):
                span = int(c.get("colspan", 1))
                cells.append((c.get_text(" ", strip=True), span))
                cells += [("", 0)] * (span - 1)
            grid.append(cells)
        ncol = max(len(r) for r in grid)
        fs = 7 if ncol >= 10 else 8
        widths = col_widths(grid, ncol, fs)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_props(t)
        for ri, (sect, tr) in enumerate(rows):
            row = t.rows[ri]
            row_flag(row, "cantSplit")
            if sect == "head":
                row_flag(row, "tblHeader")
            rcls = tr.get("class", [])
            ci = 0
            for c in tr.find_all(["td", "th"], recursive=False):
                span = int(c.get("colspan", 1))
                cell = row.cells[ci]
                if span > 1:
                    cell = cell.merge(row.cells[min(ci + span - 1, ncol - 1)])
                p = cell.paragraphs[0]
                tight(p, 0, 0, 1.05)
                ccls = c.get("class", [])
                right = "num" in ccls or "r" in ccls or "right" in c.get("style", "")
                if right:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "q" in ccls or "center" in c.get("style", ""):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if sect == "head":
                    shade(cell, NAVY)
                    inline(p, c, fs, True, "FFFFFF")
                elif "grp" in rcls:
                    shade(cell, "E4F6F8")
                    inline(p, c, fs, True, NAVY)
                elif "tot" in rcls:
                    shade(cell, "EEF2F9")
                    inline(p, c, fs, True, INK)
                elif c.name == "th":
                    shade(cell, "EEF2F9")
                    inline(p, c, fs, True, NAVY2)
                else:
                    color = TEAL if "sym" in ccls else (MUTED if ("note" in ccls or "basis" in ccls) else INK)
                    inline(p, c, fs, "nm" in ccls, color, mono="sym" in ccls)
                ci += span
        for i, col in enumerate(t.columns):      # 표 격자(gridCol) 폭 — Word·LibreOffice가 이 값으로 열을 그림
            col.width = widths[i]
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[min(i, ncol - 1)]
        tight(self.doc.add_paragraph(), 0, 4)

    def walk(self, el):
        for ch in el.children:
            if isinstance(ch, NavigableString):
                if ch.strip():
                    p = self.doc.add_paragraph()
                    set_font(p.add_run(ch.strip()), 9, color=INK)
                continue
            if not isinstance(ch, Tag):
                continue
            n, cls = ch.name, ch.get("class", [])
            if n == "section":
                if not self.first:
                    self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                self.first = False
                self.walk(ch)
            elif n == "h1":
                self.heading(ch, 20, NAVY, 0)
            elif n == "h2":
                self.heading(ch, 14, NAVY, 6)
            elif n == "h3":
                self.heading(ch, 11, NAVY2, 8)
            elif n == "p" and "eyebrow" in cls:
                p = self.doc.add_paragraph()
                tight(p, 6, 0)
                p.paragraph_format.keep_with_next = True
                inline(p, ch, 8, False, TEAL)
            elif n == "p" and "note-box" in cls:
                self.para(ch, 8.5, INK, box=True)
            elif n == "p":
                self.para(ch)
            elif n in ("ul", "ol") and ({"legend", "lg"} & set(cls)):
                self.legend(ch)
            elif n in ("ul", "ol"):
                self.bullets(ch)
            elif n == "table":
                self.table(ch)
            elif n == "svg":
                self.figure_svg(ch)
            elif n == "figcaption":
                self.para(ch, 8, MUTED, italic=True)
            elif n == "dl":
                items = [(d.find("dt"), d.find("dd")) for d in ch.find_all("div", recursive=False) if d.find("dt")]
                html = "".join(f"<tr><th>{a.decode_contents()}</th><td>{b.decode_contents()}</td></tr>" for a, b in items)
                self.table(BeautifulSoup(f"<table><tbody>{html}</tbody></table>", "html.parser").table)
            elif n == "div" and "kpis" in cls:
                cells = "".join(f"<td class='q'><b>{d.find('b').get_text(strip=True)}</b><br>{d.find('span').get_text(strip=True)}</td>"
                                for d in ch.find_all("div", recursive=False))
                self.table(BeautifulSoup(f"<table><tbody><tr>{cells}</tr></tbody></table>", "html.parser").table)
            else:
                self.walk(ch)


def docx(src, dst):
    s = open(src, encoding="utf-8").read()
    css = "".join(re.findall(r"<style>.*?</style>", s, re.S))
    soup = BeautifulSoup(s, "html.parser")
    clean(soup)
    w = Writer(css, tempfile.mkdtemp())
    if soup.find("header", class_="mast"):
        w.walk(soup.find("header", class_="mast"))
    w.walk(soup.find("main"))
    if soup.title:
        w.doc.core_properties.title = soup.title.get_text()
    w.doc.save(dst)
    return w.nfig


if __name__ == "__main__":
    if not os.path.exists(CHROME):
        sys.exit("Chrome 없음 — PDF·Word 생성 건너뜀")
    for p in PAGES:
        src = os.path.join("docs", p)
        pdf(src, src.replace(".html", ".pdf"))
        n = docx(src, src.replace(".html", ".docx"))
        print(p, "pdf", os.path.exists(src.replace(".html", ".pdf")), "docx", os.path.exists(src.replace(".html", ".docx")), f"figs={n}")
